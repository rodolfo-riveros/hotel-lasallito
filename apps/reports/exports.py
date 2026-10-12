"""Exportación genérica a CSV, Excel y PDF (RF-REP-06, RNF-DIS-03)."""

import csv
import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.core.pdf import render_pdf


def export_table(title, header, rows, fmt, filename="reporte"):
    header = [str(h) for h in header]
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        resp.write("﻿")
        w = csv.writer(resp)
        w.writerow(header)
        for row in rows:
            w.writerow(row)
        return resp
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = str(title)[:30]
        ws.append([str(title)])
        ws["A1"].font = Font(bold=True, size=13, color="7A1E2C")
        ws.append(header)
        for c in ws[2]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="F5F5F4")
        for row in rows:
            ws.append([v if isinstance(v, int | float | Decimal) else str(v) for v in row])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(45, max(10, width + 2))
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(
            buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return resp
    pdf = render_pdf("reports/table_pdf.html", {"title": title, "header": header, "rows": rows})
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return resp
